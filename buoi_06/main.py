import os 
import openpyxl 

excel_file = "user_information.xlsx"

# Kiểm tra file "user_information.xlsx" đã tồn tại chưa?
# Logic: Có file -> load_workbook()
# Logic: Chưa có file -> Workbook()
if os.path.exists(excel_file):
    workbook = openpyxl.load_workbook(excel_file) # Mở workbook đã tồn tại
    worksheet = workbook.active
else:
    # 3 bước tạo file excel: Tạo Workbook, Lấy Worksheet, Đặt tên Worksheet
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Thông tin người dùng"

    # Lưu 5 thông tin của nhân viên: Họ tên, Ngày sinh, Email, SĐT, Công việc 
    headers = [
        "Họ và tên",
        "Ngày sinh",
        "Email",
        "Số điện thoại",
        "Công việc"
    ]
    worksheet.append(headers)

# Nhập nhiều người dùng
while True:
    print("\n Nhập thông tin người dùng vào đây")

    # Nhập 5 trường thông tin
    ho_ten = input("Nhập họ tên: ")
    ngay_sinh = input("Nhập ngày sinh: ")
    email = input("Nhập Email: ")
    so_dien_thoai = input("Nhập số điện thoại: ")
    cong_viec = input("Nhập công việc:")

    # Dữ liệu người dùng
    user_data = [
        ho_ten,
        ngay_sinh,
        email,
        so_dien_thoai,
        cong_viec
    ]
    worksheet.append(user_data)
    workbook.save(excel_file)
    print("Đã lưu thông tin thành công")

    tiep_tuc = input("Bạn có muốn tiếp tục ko? (y/n): ").strip().lower()
    if tiep_tuc != "y":
        print("Kết thúc chương trình")
        break
