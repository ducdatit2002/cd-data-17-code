# Xây dựng chuyên gia xử lý hồ sơ tên CoverLetterProcessor
# Nhiệm vụ: Đọc file word, Trích xuất dữ liệu, Làm việc với Excel, Lưu kết quả 

import os # làm việc file, foler, đường dẫn
import re # xử lý regular expression 

from openpyxl import Workbook, load_workbook # Dùng để đọc và ghi excel 
from docx import Document # Xử lý file word

class CoverLetterProcessor:
    def __init__(self, folder_path, excel_file):
        self.folder_path = folder_path
        self.excel_file = excel_file

        self.wb = None
        self.ws = None

        self.headers = [
            "Họ và tên",
            "Giới tính",
            "Ngày sinh",
            "Nơi sinh",
            "Nguyên quán",
            "Hộ khẩu thường trú",
            "Chỗ ở hiện nay",
            "Điện thoại"
        ]

        self.patterns = {
            "Họ và tên": r"Họ và tên\s*:?\s*(.*?)\s*Nam/Nữ",
            "Giới tính": r"Nam/Nữ\s*:?\s*(.*?)\s*Sinh ngày",
            "Ngày sinh": r"Sinh ngày\s*:?\s*(.*?)\s*Nơi sinh",
            "Nơi sinh": r"Nơi sinh\s*:?\s*(.*?)\s*Nguyên quán",
            "Nguyên quán": r"Nguyên quán\s*:?\s*(.*?)\s*Hộ khẩu",
            "Hộ khẩu thường trú": r"Hộ khẩu.*?:?\s*(.*?)\s*Chỗ ở",
            "Chỗ ở hiện nay": r"Chỗ ở.*?:?\s*(.*?)\s*(?:Điện thoại|Số điện thoại)",
            "Điện thoại": r"(?:Điện thoại|Số điện thoại)\s*:?\s*([0-9+\s.-]+)"
        }
    # Phương thức 1: Khởi tạo, mở excel -> initialize_excel()
    def initialize_excel(self):
        if os.path.exists(self.excel_file):
            self.wb = load_workbook(self.excel_file)
            self.ws = self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active

            self.ws.title = "Thông tin người dùng"
            self.ws.append(self.headers)

    # Phương thức 2: Read Docx
    def read_docx(self, file_path):
        doc = Document(file_path)

        doc_content = [
            paragraph.text
            for paragraph in doc.paragraphs
        ]

        doc_full = "\n".join(doc_content)
        return doc_full

    # Phương thứ 3: Trích xuất thông tin 
    def extract_info(self, text):
        info = {} # tạo được 1 dictionary cụ thể

        for key in self.patterns:
            match = re.search(
                self.patterns[key],
                text,
                re.MULTILINE 
                # flag phục vụ việc xử lý văn bản nhiều dòng trong những pattern có liên quan đến đầu và cuối câu
            )
            if match: 
                info[key] = match.group(1).strip()
                # Group(): lấy phần dữ liệu được capture bởi cặp ngoặc
                # Strip: xử lý khoảng trắng đầu hoặc cuối hoặc dòng dư 
        return info

    def process_documents(self):
        self.initialize_excel()

        doc_files = os.listdir(self.folder_path) # Lấy danh sách tất cả file bên trong folder
        for file_name in doc_files:
            if not file_name.lower().endswith(".docx"):
                continue
            file_path = os.path.join(self.folder_path, file_name)

            document_text = self.read_docx(file_path)
            data = self.extract_info(document_text) # Trả ra dictionary 
            values = [
                data.get(header, "") 
                for header in self.headers
            ]
            self.ws.append(values)
        self.wb.save(self.excel_file)
        print("Xử lý thành công")        

        # {
        # "Họ và tên": "Lê Văn C",
        # "Giới tính": "Nam",
        # "Ngày sinh": "20/08/2003",
        # "Nơi sinh": "Đà Nẵng",
        # "Nguyên quán": "Quảng Nam",
        # "Hộ khẩu thường trú": "...",
        # "Chỗ ở hiện nay": "...",
        # "Điện thoại": "090..."
        # }

        # Khối 1: initialize_excel() -> mở hoặc tạo excel
        # Khối 2: read_docx() -> đọc word và trả về text
        # Khối 3: extract_info() -> nhận text -> trích xuất thông tin -> trả về dictionary

        # Pipeline:
        # B1: Khởi tạo excel
        # B2: Lấy danh sách file word
        # B3: Duyệt từng file
        # B4: Đọc file
        # B5: Trích xuất thông tin
        # B6: Đưa dữ liệu vào excel
        # B7: Lưu excel 