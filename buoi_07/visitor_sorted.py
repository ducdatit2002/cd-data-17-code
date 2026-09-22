import openpyxl

def bubble_sort(data, column):
    n = len(data)
    for i in range(n):
        for j in range(0, n - i - 1):
            if data[j][column] > data[j+1][column]:
                data[j], data[j+1] = data[j+1], data[j]
    return data

# Quy trình xử lý file excel
# B1: Đọc dữ liệu Excel
workbook = openpyxl.load_workbook("visitor_data.xlsx")
worksheet = workbook["Sheet1"]

# B2: Đưa các hàng dữ liệu vào Python
data = []

for row in worksheet.iter_rows(
    min_row = 2,
    values_only = True
):
    data.append(list(row))

# B3: Sắp xếp dữ liệu theo Pricing Plan ID
sorted_data = bubble_sort(data, 6)

# B4: Ghi dữ liệu đã sắp xếp trở lại worksheet
for row_index, row in enumerate(sorted_data, start=2):
    for column_index, value in enumerate(row):
        worksheet.cell(
            row = row_index,
            column=column_index + 1,
            value=value
        )
# B5: Lưu thành file mới 
workbook.save("visitor_data_sorted.xlsx")

# Một số lỗi thường gặp:
# 1. Sai tên file: visitor_data thay vì visitor_data.xlsx
# 2. ModuleNotFoundError: cài openpyxl ở môi trường khác, chạy ở môi trường khác
# 3. File Excel đang mở -> tắt file xong mở lại 
# 4. Sai indentation -> chú ý mối quan hệ giữa if và else 
# 5. Lỗi quên tăng row_num (quên thêm row_num += 1)
# 6. Nhầm index của Python và vị trí Excel 

# Tổng kết nội dung buổi 7:
# 1. Hiểu được cập nhật dữ liệu Excel: 
# - đọc file excel bằng openpyxl 
# - kiểm tra bằng os
# - thêm cột, đọc từng hàng
# - dùng if,else
# - cập nhật và lưu file excel 

# 2. Bubble Sort
# - so sánh 2 phần tử cạnh nhau
# - sai thứ tự thì hoán đổi
# - vì sao cần vòng lặp lồng nhau
# - sắp xếp tăng dần, giảm dần

# 3. Ứng dụng Bubble Sore vào Excel 
# - Chuyển từ "Danh sách 1 chiều" -> "Bảng dữ liệu 2 chiều"

